#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
课表解析脚本：将 .doc / .docx / .html(.htm) 课程表解析为结构化 JSON。

支持格式：
- .doc  : OLE2 二进制文档，经 Word COM 另存为 .docx 再解析（需本机装 Word，且 Python 有 pywin32）
          若文件实为 HTML/RTF（部分教务系统导出 .doc 其实是网页），自动嗅探并走对应解析
- .docx : python-docx 解析表格
- .html : 标准库 html.parser 解析（无需第三方依赖），自动识别 utf-8 / gb18030 编码

输出：**强制 UTF-8** 的标准 JSON 到 stdout，结构 { "courses": [ {...} ] }
调用：python parse_schedule.py <输入文件>
"""
import sys, os, re, json, shutil, tempfile

# ---------- 关键修复：强制 UTF-8 输出 ----------
# 背景：Windows 上 Python 子进程的 stdout 默认走 ANSI 代码页(cp936/GBK)，
# 而调用方（Node spawnSync encoding:"utf-8"）按 UTF-8 解码，
# 导致所有中文变成 U+FFFD(�)。例如 教室"徽文楼507"在 GBK 下字节为 ...C2A5 35 30 37，
# 其中「楼」= C2A5 被当作 UTF-8 解码后正好显示为 "¥"，即用户看到的 "¥507"。
# 双保险：① 重设文本层编码；② 最终结果直接用 buffer 写 UTF-8 字节，绕开文本层。
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def emit(obj):
    """把对象以 UTF-8 字节写到 stdout（绕开文本层编码，杜绝 cp936 乱码）。"""
    payload = json.dumps(obj, ensure_ascii=False).encode("utf-8")
    buf = getattr(sys.stdout, "buffer", None)
    if buf is not None:
        try:
            buf.write(payload)
            buf.write(b"\n")
            buf.flush()
            return
        except Exception:
            pass
    sys.stdout.write(payload.decode("utf-8", "replace") + "\n")


# ---------- 文件类型嗅探 ----------
def sniff_kind(path):
    """.doc 的真身可能是 OLE2 / HTML / RTF，返回 'ole' | 'html' | 'rtf'。"""
    with open(path, "rb") as f:
        head = f.read(4096)
    if head.startswith(b"\xd0\xcf\x11\xe0"):          # OLE2 复合文档
        return "ole"
    if head[:5] == b"{\\rtf":                          # RTF
        return "rtf"
    low = head.lower()
    if b"<html" in low or b"<table" in low or b"<!doctype" in low:
        return "html"
    return "ole"


# ---------- Word (.doc) 转换 ----------
def convert_doc_to_docx(doc_path):
    """用 Word COM 把 .doc 转为 .docx，返回 .docx 路径。失败抛异常。"""
    import win32com.client
    tmp = tempfile.gettempdir()
    ascii_src = os.path.join(tmp, "schedule_in.doc")
    shutil.copyfile(doc_path, ascii_src)
    out = os.path.join(tmp, "schedule_out.docx")
    if os.path.exists(out):
        os.remove(out)
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = False
    try:
        doc = word.Documents.Open(ascii_src)
        doc.SaveAs(out, FileFormat=12)  # wdFormatXMLDocument
        doc.Close()
    finally:
        word.Quit()
    return out


# ---------- 文本归一化 ----------
def normalize_period(s):
    s = (s or "").replace(" ", "").replace("\n", "").replace("\r", "").replace("\u3000", "").strip()
    if "上午" in s or s == "上":
        return "上午"
    if "下午" in s or s == "下":
        return "下午"
    if "晚上" in s or s == "晚":
        return "晚上"
    return s or None


WEEKS_RE = re.compile(r"(\d+)\s*-\s*(\d+)\s*周")
COURSE_RE = re.compile(r"^(.*?)(?:(\d+)班)?\[([^\[\]]+)\](.+?)\[([^\[\]]*)\]\s*$")
NAME_WEEKS_RE = re.compile(r"^(.*?)\[([^\[\]]+)\]\s*$")
DAY_NAMES = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]


def parse_course_cell(text):
    """解析形如 '环境数据模型与模拟1班[7-10周]易明建[徽文楼507]' 的课程单元格。"""
    text = (text or "").strip()
    blank = {"name": "", "className": None, "teacher": None, "room": None,
             "weeks": None, "startWeek": None, "endWeek": None}
    if not text:
        return blank

    m = COURSE_RE.match(text)
    if m:
        name = (m.group(1) or "").strip()
        cls = m.group(2)
        weeks = (m.group(3) or "").strip()
        teacher = (m.group(4) or "").strip()
        room = (m.group(5) or "").strip()
        if not name:  # 退路：正则吞掉了课程名
            name = text
            teacher = teacher or None
    else:
        m2 = NAME_WEEKS_RE.match(text)
        if m2:
            name = (m2.group(1) or "").strip()
            cls = None
            weeks = (m2.group(2) or "").strip()
            teacher = None
            room = None
        else:
            return {**blank, "name": text}

    sw = ew = None
    wm = WEEKS_RE.search(weeks or "")
    if wm:
        sw, ew = int(wm.group(1)), int(wm.group(2))
    else:
        single = re.search(r"(\d+)\s*周", weeks or "")
        if single:
            sw = ew = int(single.group(1))

    return {
        "name": name + (f"{cls}班" if cls else ""),
        "className": (f"{cls}班" if cls else None),
        "teacher": teacher or None,
        "room": room or None,
        "weeks": weeks or None,
        "startWeek": sw,
        "endWeek": ew,
    }


def split_cell_lines(cell_text):
    """一个单元格可能竖排多门课（换行分隔），逐条返回非空文本。"""
    return [p.strip() for p in re.split(r"[\r\n]+", cell_text or "") if p.strip()]


# ---------- 节次解析 ----------
SECTION_RANGE_RE = re.compile(r"(\d+)\s*[-~～]\s*(\d+)\s*节")
SECTION_RANGE_PLAIN_RE = re.compile(r"(\d+)\s*[-~～]\s*(\d+)")
SECTION_SINGLE_RE = re.compile(r"(\d+)\s*节")
SECTION_NUM_RE = re.compile(r"(\d+)")


def parse_section_range(text, fallback):
    """从节次单元格解析 (start, end)。
    '1-2节' -> (1,2)；'第3节' -> (3,3)；纯数字 '3' -> (3,3)。
    都解析不出时才用 fallback。
    """
    if text:
        for pat in (SECTION_RANGE_RE, SECTION_RANGE_PLAIN_RE):
            m = pat.search(text)
            if m:
                return int(m.group(1)), int(m.group(2))
        m = SECTION_SINGLE_RE.search(text) or SECTION_NUM_RE.search(text)
        if m:
            s = int(m.group(1))
            return s, s
    return fallback, fallback


# ---------- 通用行矩阵解析（docx 与 html 共用）----------
def parse_rows(rows):
    """rows: list[list[str]]，首行为表头。返回 records 列表（同课程连续节次已合并）。"""
    if not rows:
        return []
    header = [(c or "").strip().replace(" ", "").replace("\n", "") for c in rows[0]]

    day_col = {}
    for i, h in enumerate(header):
        for d, name in enumerate(DAY_NAMES, start=1):
            if name in h:
                day_col[i] = d
    if not day_col:
        return []

    section_col = None
    for i, h in enumerate(header):
        if "节次" in h or h == "节" or "节" in h:
            section_col = i
            break

    period_col = None
    for i in range(len(header)):
        if i == section_col or i in day_col:
            continue
        period_col = i
        break

    records = []
    data_ridx = 0
    for r in rows[1:]:
        if not any((c or "").strip() for c in r):
            continue  # 跳过完全空行
        data_ridx += 1
        period = normalize_period(r[period_col]) if (period_col is not None and period_col < len(r)) else None

        sec_text = (r[section_col] or "") if (section_col is not None and section_col < len(r)) else ""
        if section_col is not None and not re.search(r"\d", sec_text):
            continue  # 节次列非数字（表尾/说明行），跳过
        s_start, s_end = parse_section_range(sec_text, data_ridx)

        for ci, day in day_col.items():
            if ci >= len(r):
                continue
            cell = (r[ci] or "").strip()
            if not cell:
                continue
            for chunk in split_cell_lines(cell):
                info = parse_course_cell(chunk)
                if not info.get("name"):
                    continue
                records.append({
                    "dayOfWeek": day,
                    "sectionStart": s_start,
                    "sectionEnd": s_end,
                    "period": period,
                    "raw": chunk,
                    **info,
                })

    # 同一门课（同名/同教师/同教室/同周次）在同一星期的相邻节次合并为一个块，
    # 还原真实的"第1-2节"。注意：一个格子里可能并列多门课，因此必须按课程分组判断，
    # 不能只和"上一条记录"比较（否则多课程格会打断合并链）。
    records.sort(key=lambda x: (x["dayOfWeek"], x["sectionStart"], x["sectionEnd"]))
    buckets = {}
    for c in records:
        key = (c["dayOfWeek"], c["name"], c.get("teacher") or "", c.get("room") or "", c.get("weeks") or "")
        blocks = buckets.get(key)
        if blocks and blocks[-1]["sectionEnd"] + 1 >= c["sectionStart"]:
            if c["sectionEnd"] > blocks[-1]["sectionEnd"]:
                blocks[-1]["sectionEnd"] = c["sectionEnd"]
            continue
        if blocks is None:
            blocks = buckets.setdefault(key, [])
        blocks.append(dict(c))

    merged = [b for blist in buckets.values() for b in blist]
    merged.sort(key=lambda x: (x["dayOfWeek"], x["sectionStart"]))
    return merged


# ---------- HTML 表格解析 ----------
from html.parser import HTMLParser


class HTMLScheduleParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables = []
        self._depth = 0
        self._grid = None
        self._row = None
        self._cell = None
        self._queue = []

    def handle_starttag(self, tag, attrs):
        # 单元格内的 <br>/<p> 视为换行：教务系统常在一格里竖排多门课
        if tag in ("br", "p") and self._cell is not None:
            self._cell["text"].append("\n")
            return
        if tag == "table":
            self._depth += 1
            self._grid = []
        elif tag == "tr" and self._depth > 0:
            self._row = []
            for item in self._queue:
                col, remain, text = item
                while len(self._row) <= col:
                    self._row.append("")
                self._row[col] = text
                item[1] = remain - 1
            self._queue = [it for it in self._queue if it[1] > 0]
        elif tag in ("td", "th") and self._depth > 0:
            self._cell = {"text": [], "rs": 1, "cs": 1}
            for k, v in attrs:
                kl = (k or "").lower()
                if kl == "rowspan":
                    try:
                        self._cell["rs"] = max(1, int(v))
                    except Exception:
                        pass
                elif kl == "colspan":
                    try:
                        self._cell["cs"] = max(1, int(v))
                    except Exception:
                        pass

    def handle_data(self, data):
        if self._cell is not None:
            self._cell["text"].append(data)

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None and self._depth > 0:
            text = "".join(self._cell["text"]).strip()
            cs = self._cell["cs"]
            rs = self._cell["rs"]
            start = len(self._row)
            for i in range(cs):
                while len(self._row) < start + i + 1:
                    self._row.append("")
                self._row[start + i] = text if i == 0 else ""
            if rs > 1:
                self._queue.append([start, rs - 1, text])
            self._cell = None
        elif tag == "tr" and self._depth > 0 and self._row is not None:
            self._grid.append(self._row)
            self._row = None
        elif tag == "table":
            if self._depth == 1 and self._grid is not None:
                self.tables.append(self._grid)
            self._depth -= 1
            self._grid = None


def read_text_auto(path):
    """自动识别编码读取文本（教务系统导出的 HTML 多为 GBK/GB2312）。"""
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "utf-8"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    for enc in ("gb18030", "big5", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="ignore")


def parse_html_file(path):
    html = read_text_auto(path)
    parser = HTMLScheduleParser()
    parser.feed(html)
    courses = []
    for grid in parser.tables:
        courses.extend(parse_rows(grid))
    return courses


# ---------- docx 解析 ----------
def parse_docx_file(path):
    import docx
    document = docx.Document(path)
    if not document.tables:
        return None  # 调用方处理“未找到表格”
    courses = []
    for t in document.tables:
        rows = [[(c.text or "").strip() for c in row.cells] for row in t.rows]
        courses.extend(parse_rows(rows))
    return courses


# ---------- 主流程 ----------
def main():
    if len(sys.argv) < 2:
        emit({"error": "缺少输入文件参数"})
        sys.exit(1)
    in_path = sys.argv[1]
    if not os.path.exists(in_path):
        emit({"error": f"文件不存在: {in_path}"})
        sys.exit(1)

    ext = os.path.splitext(in_path)[1].lower()
    docx_path = None
    try:
        kind = sniff_kind(in_path)

        # 真身是 HTML/RTF 的 .doc（教务系统常见）直接按 HTML 解析
        if ext in (".html", ".htm") or (ext == ".doc" and kind == "html"):
            courses = parse_html_file(in_path)
            if not courses and ext == ".doc":
                docx_path = convert_doc_to_docx(in_path)
                courses = parse_docx_file(docx_path)
        elif ext == ".doc" and kind == "rtf":
            # RTF：借 Word 转 docx
            docx_path = convert_doc_to_docx(in_path)
            courses = parse_docx_file(docx_path)
        else:
            target = in_path
            if ext == ".doc":
                docx_path = convert_doc_to_docx(in_path)
                target = docx_path
            courses = parse_docx_file(target)

        if courses is None:
            emit({"error": "文档中未找到可识别的课表，请确认表格包含「节次」与「星期」表头"})
            sys.exit(1)

        # 去重（极少数重复行）
        seen = set()
        uniq = []
        for c in courses:
            key = (c["name"], c["dayOfWeek"], c["sectionStart"], c["sectionEnd"])
            if key in seen:
                continue
            seen.add(key)
            uniq.append(c)

        emit({"courses": uniq})
    except Exception as e:
        emit({"error": f"解析失败: {e}"})
        sys.exit(1)
    finally:
        if docx_path and os.path.exists(docx_path):
            try:
                os.remove(docx_path)
            except Exception:
                pass


if __name__ == "__main__":
    main()
